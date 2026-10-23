"""
DOCX Parser implementation using python-docx and OpenXML inspection.
Extracts paragraphs, tables, and inline drawing images in exact document order.
Anchors screenshots directly below corresponding instruction steps.
Extracts embedded OLE object packages recursively.
"""

from __future__ import annotations

import hashlib
import io
import logging
import zipfile
from typing import Any, Dict, List, Optional, Set

from ...domain.entities import ParsedDocument, ParsedPage
from ...domain.exceptions import DocumentParsingError
from .base import BaseDocumentParser

logger = logging.getLogger("aip-rag.parser.docx")


class DocxDocumentParser(BaseDocumentParser):
    """
    Parser for Microsoft Word (.docx) documents with full table and multimodal image alignment support.
    Extracts inline screenshots from OpenXML drawing markup and embeds Markdown image links
    directly beneath the corresponding paragraph step.
    """

    def __init__(
        self,
        base_storage_url: str = "https://storage.enterprise.internal/media",
    ):
        self.base_storage_url = base_storage_url.rstrip("/")

    @staticmethod
    def _table_to_markdown(table) -> str:
        """Convert a python-docx Table into clean Markdown table format."""
        try:
            rows_data = []
            for row in table.rows:
                row_cells = []
                seen_tc = set()
                for cell in row.cells:
                    tc_id = id(cell._tc)
                    if tc_id in seen_tc:
                        row_cells.append("")
                        continue
                    seen_tc.add(tc_id)
                    txt = " ".join(cell.text.split()).replace("|", "\\|")
                    row_cells.append(txt)
                if any(bool(c.strip()) for c in row_cells):
                    rows_data.append(row_cells)

            if not rows_data:
                return ""

            num_cols = max(len(r) for r in rows_data)
            header = rows_data[0]
            while len(header) < num_cols:
                header.append("")

            lines = [
                "| " + " | ".join(header) + " |",
                "| " + " | ".join(["---"] * num_cols) + " |",
            ]
            for r in rows_data[1:]:
                while len(r) < num_cols:
                    r.append("")
                lines.append("| " + " | ".join(r[:num_cols]) + " |")

            return "\n".join(lines)
        except Exception as exc:
            logger.warning("Error converting Word table to Markdown: %s", exc)
            return ""

    @staticmethod
    def _format_paragraph(p) -> str:
        """Format a Word paragraph, prefixing Markdown headings if style indicates a heading."""
        txt = p.text.strip()
        if not txt:
            return ""
        style_name = (getattr(getattr(p, "style", None), "name", "") or "").strip().lower()
        if style_name in ("title", "heading 1"):
            return f"# {txt}"
        elif style_name in ("subtitle", "heading 2"):
            return f"## {txt}"
        elif style_name == "heading 3":
            return f"### {txt}"
        elif style_name == "heading 4":
            return f"#### {txt}"
        elif style_name == "heading 5":
            return f"##### {txt}"
        elif style_name == "heading 6":
            return f"###### {txt}"
        return txt

    def _extract_images_from_paragraph(
        self,
        p,
        doc,
        seen_hashes: Set[str],
        all_images: List[Dict[str, Any]],
        step_label: Optional[str] = None,
    ) -> List[str]:
        """
        Extracts inline DrawingML images from a paragraph element via OpenXML relationship mapping.
        Returns markdown image tags to append right below the paragraph.
        """
        image_tags: List[str] = []
        try:
            # 1. Standard DrawingML blip images (<a:blip r:embed="rIdX">)
            blip_nodes = p._p.xpath(".//a:blip")
            for blip in blip_nodes:
                r_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if r_id and r_id in doc.part.related_parts:
                    img_part = doc.part.related_parts[r_id]
                    img_bytes = getattr(img_part, "blob", b"")
                    if not img_bytes or len(img_bytes) < 100:
                        continue

                    # Extract extension from content_type (e.g. image/png -> png)
                    content_type = getattr(img_part, "content_type", "image/png")
                    ext = content_type.split("/")[-1].replace("jpeg", "jpg") if "/" in content_type else "png"

                    img_hash = hashlib.sha256(img_bytes).hexdigest()[:16]
                    img_url = f"{self.base_storage_url}/{img_hash}.{ext}"

                    if img_hash not in seen_hashes:
                        seen_hashes.add(img_hash)
                        all_images.append({
                            "hash": img_hash,
                            "url": img_url,
                            "ext": ext,
                            "size_bytes": len(img_bytes),
                            "data": img_bytes,
                        })

                    caption = step_label or "Ảnh minh họa"
                    image_tags.append(f"![{caption}]({img_url})")

            # 2. Legacy VML imagedata (<v:imagedata r:id="rIdX">)
            vml_nodes = p._p.xpath(".//v:imagedata")
            for vml in vml_nodes:
                r_id = vml.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                if r_id and r_id in doc.part.related_parts:
                    img_part = doc.part.related_parts[r_id]
                    img_bytes = getattr(img_part, "blob", b"")
                    if not img_bytes or len(img_bytes) < 100:
                        continue

                    ext = "png"
                    img_hash = hashlib.sha256(img_bytes).hexdigest()[:16]
                    img_url = f"{self.base_storage_url}/{img_hash}.{ext}"

                    if img_hash not in seen_hashes:
                        seen_hashes.add(img_hash)
                        all_images.append({
                            "hash": img_hash,
                            "url": img_url,
                            "ext": ext,
                            "size_bytes": len(img_bytes),
                            "data": img_bytes,
                        })

                    caption = step_label or "Ảnh minh họa"
                    image_tags.append(f"![{caption}]({img_url})")

        except Exception as exc:
            logger.debug("Non-fatal error inspecting paragraph drawing: %s", exc)

        return image_tags

    def parse(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        try:
            import docx
        except ImportError as exc:
            raise DocumentParsingError("python-docx is required for DOCX parsing") from exc

        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            blocks: List[str] = []
            all_extracted_images: List[Dict[str, Any]] = []
            seen_image_hashes: Set[str] = set()

            # 1. Iterate document body children in natural order (Paragraphs + Tables)
            try:
                body = doc.element.body
                for child in body.iterchildren():
                    if isinstance(child, docx.oxml.text.paragraph.CT_P):
                        p = docx.text.paragraph.Paragraph(child, doc)
                        formatted_p = self._format_paragraph(p)

                        # Determine if this paragraph represents a step (e.g. Bước 1, Bước 2...)
                        step_label = None
                        if formatted_p:
                            for prefix in ("bước", "step", "thao tác"):
                                if formatted_p.lower().startswith(prefix):
                                    step_label = formatted_p.split(":")[0].strip("#* ")
                                    break

                        img_tags = self._extract_images_from_paragraph(
                            p, doc, seen_image_hashes, all_extracted_images, step_label=step_label
                        )

                        if formatted_p:
                            blocks.append(formatted_p)

                        # Anchor image tags immediately following paragraph text
                        for tag in img_tags:
                            blocks.append(tag)

                    elif isinstance(child, docx.oxml.table.CT_Tbl):
                        t = docx.table.Table(child, doc)
                        tbl_md = self._table_to_markdown(t)
                        if tbl_md:
                            blocks.append(tbl_md)
            except Exception as iter_exc:
                logger.debug("Falling back to paragraphs + tables sequential extraction: %s", iter_exc)
                for p in doc.paragraphs:
                    formatted_p = self._format_paragraph(p)
                    img_tags = self._extract_images_from_paragraph(
                        p, doc, seen_image_hashes, all_extracted_images
                    )
                    if formatted_p:
                        blocks.append(formatted_p)
                    for tag in img_tags:
                        blocks.append(tag)

                for t in doc.tables:
                    tbl_md = self._table_to_markdown(t)
                    if tbl_md:
                        blocks.append(tbl_md)

            # 2. Extract embedded OLE packages (word/embeddings/) if present in zip
            embedded_files_count = 0
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                    for name in zf.namelist():
                        if name.startswith("word/embeddings/") and not name.endswith("/"):
                            emb_bytes = zf.read(name)
                            emb_filename = name.split("/")[-1]
                            if emb_filename.lower().endswith(".docx"):
                                sub_parser = DocxDocumentParser(base_storage_url=self.base_storage_url)
                                sub_doc = sub_parser.parse(emb_bytes, filename=f"{filename}::{emb_filename}")
                                for sp in sub_doc.pages:
                                    blocks.append(f"[Đính kèm: {emb_filename}]\n{sp.text}")
                                all_extracted_images.extend(sub_doc.metadata.get("images", []))
                                embedded_files_count += 1
            except Exception as zip_exc:
                logger.debug("No zip embeddings found: %s", zip_exc)

            # 3. Group into synthetic pages of ~500 words
            pages: List[ParsedPage] = []
            current_page_text: List[str] = []
            current_words = 0
            page_num = 1

            for b in blocks:
                current_page_text.append(b)
                current_words += len(b.split())
                if current_words >= 500:
                    pages.append(ParsedPage(page_number=page_num, text="\n\n".join(current_page_text)))
                    page_num += 1
                    current_page_text = []
                    current_words = 0

            if current_page_text:
                pages.append(ParsedPage(page_number=page_num, text="\n\n".join(current_page_text)))

            logger.info(
                "Parsed DOCX '%s': %d blocks across %d pages, %d images extracted & anchored, %d embedded files",
                filename, len(blocks), len(pages), len(all_extracted_images), embedded_files_count
            )
            return ParsedDocument(
                filename=filename,
                file_type="docx",
                pages=pages,
                metadata={
                    "total_pages": len(pages),
                    "total_blocks": len(blocks),
                    "images_count": len(all_extracted_images),
                    "images": all_extracted_images,
                    "embedded_files_count": embedded_files_count,
                },
            )
        except Exception as exc:
            logger.error("Failed to parse DOCX document '%s': %s", filename, exc)
            raise DocumentParsingError(f"DOCX parsing failed: {exc}") from exc
