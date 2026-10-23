"""
Automated Unit & Integration Test Suite for FR-ING-05 & ING-RAG-05:
Step-by-Step Text-Image Alignment & Embedded Object Extraction.
Verifies compliance with BRD Chapter 4.1 & SRS Section 3 (Nhóm 1).
"""

import io
import pymupdf as fitz
import docx

from solutions.rag.infrastructure.parsers.docx_parser import DocxDocumentParser
from solutions.rag.infrastructure.parsers.pdf_parser import PdfDocumentParser
from solutions.rag.application.components.chunker import DocumentChunker


def _generate_sample_png_bytes(color: int = 128) -> bytes:
    """Generate minimal valid PNG bytes using PyMuPDF pixmap."""
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 80, 80), False)
    pix.set_rect(fitz.IRect(0, 0, 80, 80), (color, color, color))
    png_data = pix.tobytes("png")
    return png_data


def test_docx_step_by_step_image_alignment():
    """
    ING-RAG-05 DoD 1 & 2:
    In Word (.docx), inline screenshots extracted via OpenXML must appear
    immediately beneath the corresponding instruction step, never detached.
    """
    doc = docx.Document()
    doc.add_heading("Hướng Dẫn Cài Đặt Plugin Ký Số", level=1)

    # Step 1
    p1 = doc.add_paragraph("Bước 1: Tải bộ cài đặt Plugin_Setup.exe từ cổng thông tin.")
    img1_bytes = _generate_sample_png_bytes(color=50)
    p1.add_run().add_picture(io.BytesIO(img1_bytes), width=docx.shared.Inches(2.0))

    # Step 2
    p2 = doc.add_paragraph("Bước 2: Mở trình duyệt Chrome và bật tiện ích ký số SmartCA.")
    img2_bytes = _generate_sample_png_bytes(color=200)
    p2.add_run().add_picture(io.BytesIO(img2_bytes), width=docx.shared.Inches(2.0))

    # Concluding step without image
    doc.add_paragraph("Bước 3: Xác nhận biểu tượng màu xanh xuất hiện trên thanh công cụ.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    parser = DocxDocumentParser(base_storage_url="https://storage.enterprise.internal/media")
    parsed_doc = parser.parse(docx_bytes, filename="HDSD_Plugin.docx")

    assert parsed_doc.file_type == "docx"
    assert parsed_doc.metadata.get("images_count") == 2
    assert len(parsed_doc.pages) >= 1

    full_text = parsed_doc.full_text()

    # Step 1 text must appear before Image 1
    pos_step1 = full_text.find("Bước 1: Tải bộ cài đặt")
    pos_img1 = full_text.find("![Bước 1](https://storage.enterprise.internal/media/")
    assert pos_step1 != -1, "Step 1 text missing"
    assert pos_img1 != -1, "Image 1 tag missing"
    assert pos_step1 < pos_img1, "Step 1 text must appear BEFORE Image 1"

    # Image 1 must appear before Step 2
    pos_step2 = full_text.find("Bước 2: Mở trình duyệt Chrome")
    pos_img2 = full_text.find("![Bước 2](https://storage.enterprise.internal/media/")
    assert pos_step2 != -1, "Step 2 text missing"
    assert pos_img2 != -1, "Image 2 tag missing"
    assert pos_img1 < pos_step2, "Image 1 must appear BEFORE Step 2"
    assert pos_step2 < pos_img2, "Step 2 text must appear BEFORE Image 2"

    # Step 3 must appear after Image 2
    pos_step3 = full_text.find("Bước 3: Xác nhận biểu tượng")
    assert pos_img2 < pos_step3, "Image 2 must appear BEFORE Step 3"


def test_docx_image_deduplication():
    """
    ING-RAG-05 DoD:
    Identical images (same SHA-256) used in multiple places must reuse the same URL.
    """
    doc = docx.Document()
    same_img_bytes = _generate_sample_png_bytes(color=100)

    p1 = doc.add_paragraph("Bước 1: Nhấn nút Lưu dữ liệu.")
    p1.add_run().add_picture(io.BytesIO(same_img_bytes), width=docx.shared.Inches(1.0))

    p2 = doc.add_paragraph("Bước 2: Tiếp tục nhấn nút Lưu dữ liệu để xác nhận.")
    p2.add_run().add_picture(io.BytesIO(same_img_bytes), width=docx.shared.Inches(1.0))

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    parser = DocxDocumentParser(base_storage_url="https://storage.enterprise.internal/media")
    parsed_doc = parser.parse(docx_bytes, filename="duplicate_test.docx")

    # Deduped to 1 unique image in metadata
    assert parsed_doc.metadata.get("images_count") == 1
    unique_hash = parsed_doc.metadata["images"][0]["hash"]
    expected_url = f"https://storage.enterprise.internal/media/{unique_hash}.png"

    full_text = parsed_doc.full_text()
    assert full_text.count(expected_url) == 2


def test_pdf_2d_spatial_image_alignment():
    """
    ING-RAG-05 DoD 2:
    In PDF, PyMuPDF (fitz) spatial 2D coordinates (y-axis) must sort
    text blocks and image blocks in exact vertical reading order.
    """
    pdf_doc = fitz.open()
    page = pdf_doc.new_page(width=595, height=842)  # A4

    # Insert Step 1 at top (y = 50)
    page.insert_text(fitz.Point(50, 50), "Buoc 1: Mo ung dung Hoa don dien tu", fontsize=12)

    # Insert Image 1 at y = 80 -> 180
    img1_bytes = _generate_sample_png_bytes(color=60)
    img1_rect = fitz.Rect(50, 80, 300, 180)
    page.insert_image(img1_rect, stream=img1_bytes)

    # Insert Step 2 below Image 1 (y = 220)
    page.insert_text(fitz.Point(50, 220), "Buoc 2: Chon menu Lap hoa don dieu chinh", fontsize=12)

    # Insert Image 2 below Step 2 (y = 250 -> 350)
    img2_bytes = _generate_sample_png_bytes(color=180)
    img2_rect = fitz.Rect(50, 250, 300, 350)
    page.insert_image(img2_rect, stream=img2_bytes)

    # Insert Step 3 at bottom (y = 400)
    page.insert_text(fitz.Point(50, 400), "Buoc 3: Nhan Ky dien tu va gui co quan thue", fontsize=12)

    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    parser = PdfDocumentParser(base_storage_url="https://storage.enterprise.internal/media")
    parsed_doc = parser.parse(pdf_bytes, filename="HDSD_HDDT.pdf")

    assert parsed_doc.file_type == "pdf"
    assert parsed_doc.metadata.get("images_count") == 2

    full_text = parsed_doc.full_text()

    pos_step1 = full_text.find("Buoc 1: Mo ung dung Hoa don dien tu")
    pos_img1 = full_text.find("![Ảnh minh họa Trang 1](https://storage.enterprise.internal/media/")
    pos_step2 = full_text.find("Buoc 2: Chon menu Lap hoa don dieu chinh")

    assert pos_step1 != -1, "Step 1 text missing in PDF"
    assert pos_img1 != -1, "Image 1 tag missing in PDF"
    assert pos_step2 != -1, "Step 2 text missing in PDF"

    # Strict spatial sequence: Step 1 -> Image 1 -> Step 2
    assert pos_step1 < pos_img1 < pos_step2, (
        f"PDF 2D spatial sequence failed: step1={pos_step1}, img1={pos_img1}, step2={pos_step2}"
    )


def test_chunker_preserves_multimodal_image_anchors():
    """
    ING-RAG-05 DoD 3:
    When DocumentChunker chunks a parsed document containing image anchors,
    the image tags must be preserved intact within the corresponding step chunk.
    """
    doc_text = (
        "## QUY TRÌNH KÝ SỐ VĂN BẢN\n\n"
        "Bước 1: Cắm USB Token vào máy tính và kiểm tra nhận diện.\n\n"
        "![Ảnh Bước 1](https://storage.enterprise.internal/media/usb_step1.png)\n\n"
        "Bước 2: Chọn văn bản cần ký và nhấn nút Xác nhận ký số.\n\n"
        "![Ảnh Bước 2](https://storage.enterprise.internal/media/sign_step2.png)"
    )
    from solutions.rag.domain.entities import ParsedDocument, ParsedPage
    parsed_doc = ParsedDocument(
        filename="QuyTrinh_KySo.docx",
        file_type="docx",
        pages=[ParsedPage(page_number=1, text=doc_text)],
    )

    chunker = DocumentChunker(child_chunk_size_chars=140, inject_breadcrumbs=True)
    chunks = chunker.chunk_document(parsed_doc, doc_id="doc_kyso_01")

    assert len(chunks) >= 2

    # Check that Step 1 chunk contains image anchor for step 1
    chunk_step1 = next(c for c in chunks if "Bước 1" in c.text)
    assert "![Ảnh Bước 1](https://storage.enterprise.internal/media/usb_step1.png)" in chunk_step1.text
    assert "[Ngữ cảnh: " in chunk_step1.text

    # Check that Step 2 chunk contains image anchor for step 2
    chunk_step2 = next(c for c in chunks if "Bước 2" in c.text)
    assert "![Ảnh Bước 2](https://storage.enterprise.internal/media/sign_step2.png)" in chunk_step2.text
